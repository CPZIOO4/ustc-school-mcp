import asyncio
from io import BytesIO
import json
import os
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

from school_mcp.desktop import materials, artifacts, document_io
from school_mcp.desktop.chat import ChatRunner
from school_mcp.desktop.storage import Data, write_private
from school_mcp.desktop.imports import reuse_mail
from school_mcp.bb.identity import save_session_identity
from school_mcp.service_errors import ServiceError


@unittest.skipUnless(os.name == 'nt', 'Windows encrypted profile')
class DocumentTests(unittest.TestCase):
    def setUp(self):
        env=patch.dict(os.environ); env.start(); self.addCleanup(env.stop)
        tmp=tempfile.TemporaryDirectory(); self.addCleanup(tmp.cleanup)
        self.data=Data(tmp.name); self.chat=self.data.new_chat()
        self.runner=ChatRunner(self.data, threading.Event())

    def upload(self,name,raw):
        path=self.data.profile/'uploads'/name; path.parent.mkdir(exist_ok=True); path.write_bytes(raw)
        self.chat['allowed_files'].append(str(path)); return str(path)

    def call(self,name,**args):
        return asyncio.run(self.runner.dispatch(self.chat,name,args))

    def test_generated_formats_readable_and_downloadable_only_in_same_chat(self):
        for suffix in ['txt','md','json','csv','xlsx','docx','pdf']:
            with self.subTest(suffix=suffix):
                content='{"课程":"A123"}' if suffix=='json' else '课程实验 A123\n材料准备 B456'
                result=self.call('school_local_create_file',filename='结果.'+suffix,content=content,rows=[['项目','编号'],['课程','A123']])
                self.data.save_chat(self.chat)
                raw=artifacts.download(self.data,self.chat['id'],result['artifact_id'])
                self.assertEqual(raw,Path(result['path']).read_bytes())
                read=self.call('school_local_read_file',path=result['path'])
                self.assertEqual(read['status'],'read')
                self.assertIn('A123',read['content'])
                with self.assertRaises(ValueError):
                    artifacts.download(self.data,self.data.new_chat()['id'],result['artifact_id'])

    def test_file_names_never_escape_output_or_overwrite(self):
        for filename in ['../outside.txt','C:/outside.txt','CON.txt','a.py','a.html','a.txt.','a\\b.txt']:
            with self.subTest(filename=filename), self.assertRaises(ServiceError):
                self.call('school_local_create_file',filename=filename,content='x')
        a=self.call('school_local_create_file',filename='材料.txt',content='first')
        b=self.call('school_local_create_file',filename='材料.txt',content='second')
        self.assertNotEqual(a['path'],b['path'])
        self.assertEqual(Path(a['path']).read_text(encoding='utf-8'),'first')

    def test_redirected_output_directory_and_tamper_detection(self):
        target = self.data.root / 'redirected-output'; target.mkdir()
        link = self.data.profile / 'generated'
        try:
            link.symlink_to(target, target_is_directory=True)
        except OSError:
            self.skipTest('Directory symlink unavailable on this host')
        result = self.call('school_local_create_file', filename='result.txt', content='known result')
        self.data.save_chat(self.chat)
        self.assertEqual(artifacts.download(self.data, self.chat['id'], result['artifact_id']), b'known result')
        Path(result['path']).write_text('changed', encoding='utf-8')
        with self.assertRaises(ValueError):
            artifacts.download(self.data, self.chat['id'], result['artifact_id'])

    def test_excel_sheet_rows_and_formulas_not_recalculated(self):
        from openpyxl import Workbook, load_workbook
        wb=Workbook(); ws=wb.active; ws.title='第一张'; ws.append(['编号','结果']); ws.append(['A123','=1+2']); wb.create_sheet('另一张').append(['第二表'])
        raw=BytesIO(); wb.save(raw); wb.close()
        path=self.upload('多表.xlsx',raw.getvalue())
        first=self.call('school_local_read_file',path=path,rows=1)
        self.assertEqual(first['next_row'],2); self.assertFalse(first['complete'])
        second=self.call('school_local_read_file',path=path,row=2)
        self.assertEqual(second['formulas'][0]['formula'],'=1+2')
        self.assertIsNone(second['formulas'][0]['cached_value'])
        self.assertFalse(second['formula_recalculated'])
        self.assertEqual(self.call('school_local_read_file',path=path,sheet='不存在')['status'],'needs_sheet')
        output=self.call('school_local_create_file',filename='文本公式.xlsx',rows=[['=1+2']])
        check=load_workbook(output['path']); self.assertEqual(check.active['A1'].data_type,'s'); check.close()

    def test_pdf_page_selection_and_explicit_invalid_range(self):
        from reportlab.pdfgen import canvas
        buf=BytesIO(); c=canvas.Canvas(buf)
        for n in range(1,4): c.drawString(30,800,f'PAGE {n} SAMPLE A{n}123'); c.showPage()
        c.save(); path=self.upload('pages.pdf',buf.getvalue())
        first=self.call('school_local_read_file',path=path,page=1,pages=1)
        self.assertIn('A1123',first['content']); self.assertNotIn('A2123',first['content'])
        self.assertEqual(first['next_page'],2); self.assertFalse(first['complete'])
        all_pages=self.call('school_local_read_file',path=path,pages=3)
        self.assertTrue(all_pages['complete']); self.assertIn('A3123',all_pages['content'])
        self.assertEqual(self.call('school_local_read_file',path=path,page=9)['status'],'unreadable_file')

    def test_real_local_ocr_image_and_scanned_pdf(self):
        from PIL import Image, ImageDraw, ImageFont
        from reportlab.pdfgen import canvas
        from reportlab.lib.utils import ImageReader
        image=Image.new('RGB',(1100,180),'white')
        ImageDraw.Draw(image).text((35,45),'COURSE 12345',font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',60),fill='black')
        png=BytesIO(); image.save(png,format='PNG')
        path=self.upload('scan.png',png.getvalue())
        result=self.call('school_local_read_file',path=path)
        self.assertIn('12345',result['content']); self.assertTrue(result['ocr'])
        pdf=BytesIO(); c=canvas.Canvas(pdf); c.drawImage(ImageReader(image),20,600,width=550,height=90); c.save()
        result=self.call('school_local_read_file',path=self.upload('scan.pdf',pdf.getvalue()))
        self.assertIn('12345',result['content']); self.assertTrue(result['pages'][0]['ocr'])

    def test_source_mail_checked_before_migration_and_identity_guard(self):
        save_session_identity('synthetic-student')
        source=self.data.root/'source'; source.mkdir()
        write_private(source/'ustc-identity.credentials.dpapi',{'username':'someone-else'})
        with self.assertRaises(ValueError): reuse_mail(self.data,source)
        write_private(source/'ustc-identity.credentials.dpapi',{'username':'synthetic-student'})
        (source/'mail.json').write_text(json.dumps({'address':'synthetic@mail.ustc.edu.cn'}))
        write_private(source/'mail.credentials.dpapi',{'address':'synthetic@mail.ustc.edu.cn','password':'synthetic-password'})
        with patch('school_mcp.desktop.imports.MailClient') as client:
            client.return_value.check.return_value={'connected':False}
            with self.assertRaises(ValueError): reuse_mail(self.data,source)
            self.assertFalse((self.data.profile/'mail.json').exists())
            client.return_value.check.return_value={'connected':True,'read_only':True}
            result=reuse_mail(self.data,source)
            self.assertTrue(result['reused']); self.assertFalse(result['created_client_password'])
        self.assertNotIn(b'synthetic-password',(self.data.profile/'mail.credentials.dpapi').read_bytes())


if __name__=='__main__': unittest.main()
