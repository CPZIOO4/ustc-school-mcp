"""Bounded local document extraction and generation; no macros or external links."""
from io import BytesIO, StringIO
import csv
from datetime import date, datetime
import json
import math
import threading
from zipfile import ZipFile

LOCK = threading.RLock()  # PDFium calls must not overlap across threads.
OCR = None
IMAGE_SUFFIXES = {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tif', '.tiff'}
READ_SUFFIXES = IMAGE_SUFFIXES | {'.pdf', '.xlsx'}
WRITE_SUFFIXES = {'.txt', '.md', '.csv', '.json', '.docx', '.xlsx', '.pdf'}


def ocr_image(image):
    global OCR
    import numpy as np
    from rapidocr_onnxruntime import RapidOCR
    with LOCK:
        if OCR is None:
            OCR = RapidOCR(intra_op_num_threads=2, inter_op_num_threads=1, print_verbose=False)
        image = image.convert('RGB')
        original = image.size
        image.thumbnail((3000, 3000))
        lines, _ = OCR(np.asarray(image))
        return {'text': '\n'.join(str(row[1]) for row in lines or []),
                'minimum_confidence': min((float(row[2]) for row in lines or []), default=None),
                'ocr': True, 'resized': image.size != original,
                'note': 'OCR 可能误识别数字、公式和表格关系；重要字段需与原图核对。'}


def image_text(raw):
    from PIL import Image, ImageOps
    with Image.open(BytesIO(raw)) as source:
        if source.width * source.height > 25_000_000:
            raise ValueError('image dimensions too large')
        if getattr(source, 'n_frames', 1) > 1:
            raise ValueError('multiple image frames; split the file first')
        result = ocr_image(ImageOps.exif_transpose(source))
    return {'content': result.pop('text'), 'extraction': 'local_ocr', 'document_complete': True, **result}


def pdf_text(raw, start=1, count=3, ocr='auto'):
    import pypdfium2 as pdfium
    with LOCK, pdfium.PdfDocument(raw) as doc:
        total = len(doc)
        if not 1 <= start <= total:
            raise ValueError('page out of range')
        pages = []
        end = min(start - 1 + count, total)
        for index in range(start - 1, end):
            page = doc[index]
            try:
                textpage = page.get_textpage()
                try:
                    text = textpage.get_text_range()
                finally:
                    textpage.close()
                item = {'page': index + 1, 'ocr': False}
                if ocr == 'always' or ocr == 'auto' and not text.strip():
                    width, height = page.get_size()
                    scale = min(2.5, 3000 / max(width, height))
                    bitmap = page.render(scale=scale)
                    try:
                        result = ocr_image(bitmap.to_pil())
                    finally:
                        bitmap.close()
                    text = result.pop('text')
                    item.update(result)
                elif not text.strip():
                    item['note'] = '此页未提取到文字；可使用 ocr=always 尝试扫描识别。'
                item['text'] = text
                pages.append(item)
            finally:
                page.close()
        return {'content': '\n\n'.join(f'[第 {p["page"]} 页]\n{p["text"]}' for p in pages),
                'pages': [{k:v for k,v in p.items() if k != 'text'} for p in pages],
                'page_count': total, 'next_page': end + 1 if end < total else None,
                'document_complete': start == 1 and end == total, 'extraction': 'pdf_text_and_local_ocr'}


def checked_zip(raw):
    with ZipFile(BytesIO(raw)) as archive:
        if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 40 * 1024 * 1024:
            raise ValueError('expanded document too large')


def excel_text(raw, sheet='', row=1, rows=50):
    from openpyxl import load_workbook
    checked_zip(raw)
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    values = load_workbook(BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    try:
        names = workbook.sheetnames
        if sheet and sheet not in names:
            return {'status': 'needs_sheet', 'sheets': names, 'next_step': '请选择返回的真实工作表名称。'}
        name = sheet or names[0]
        ws, cached = workbook[name], values[name]
        if ws.max_column > 200 or ws.max_row > 200000:
            raise ValueError('sheet dimensions too large')
        end = min(row + rows - 1, ws.max_row)
        if row > ws.max_row:
            raise ValueError('row out of range')
        lines, formulas = [], []
        for cells, cached_cells in zip(ws.iter_rows(min_row=row, max_row=end), cached.iter_rows(min_row=row, max_row=end)):
            parts = []
            for cell, value in zip(cells, cached_cells):
                if cell.value is None:
                    continue
                display = cell.value
                if cell.data_type == 'f':
                    cached_value = value.value.isoformat() if isinstance(value.value, (datetime, date)) else value.value
                    formulas.append({'cell': cell.coordinate, 'formula': cell.value, 'cached_value': cached_value})
                    display = f'公式 {cell.value}；缓存值 {value.value!s}（未重新计算）'
                parts.append(f'{cell.coordinate}={display}')
            lines.append('\t'.join(parts))
        return {'content': '\n'.join(lines), 'sheets': names, 'sheet': name, 'total_rows': ws.max_row,
                'total_columns': ws.max_column, 'row_start': row, 'row_end': end,
                'next_row': end + 1 if end < ws.max_row else None,
                'formulas': formulas, 'formula_recalculated': False,
                'document_complete': len(names) == 1 and row == 1 and end == ws.max_row,
                'extraction': 'xlsx_cells_and_cached_formulas'}
    finally:
        workbook.close(); values.close()


def extract(raw, suffix, args):
    if suffix == '.pdf':
        return pdf_text(raw, args.get('page', 1), args.get('pages', 3), args.get('ocr', 'auto'))
    if suffix == '.xlsx':
        return excel_text(raw, args.get('sheet', ''), args.get('row', 1), args.get('rows', 50))
    return image_text(raw)


def create(suffix, content, rows):
    buffer = BytesIO()
    if suffix in {'.txt', '.md'}:
        return content.encode('utf-8')
    if suffix == '.json':
        return (json.dumps(json.loads(content), ensure_ascii=False, indent=2, allow_nan=False) + '\n').encode('utf-8')
    if suffix == '.csv':
        stream = StringIO(newline=''); writer = csv.writer(stream)
        for row in rows:
            # Keep supplied text literal when opened in spreadsheet applications.
            writer.writerow(["'" + x if isinstance(x,str) and x.lstrip().startswith(('=', '+', '-', '@')) else x for x in row])
        return stream.getvalue().encode('utf-8-sig')
    if suffix == '.xlsx':
        from openpyxl import Workbook
        workbook = Workbook(); ws = workbook.active; ws.title = '内容'
        for line in rows:
            ws.append(line)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Model output cannot introduce formulas/macros.
        workbook.save(buffer); workbook.close()
    elif suffix == '.docx':
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Pt
        doc = Document(); style = doc.styles['Normal']; style.font.name = 'Microsoft YaHei'; style.font.size = Pt(11)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts(); fonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        for line in content.splitlines():
            doc.add_paragraph(line)
        if rows:
            table = doc.add_table(rows=0, cols=max(map(len,rows))); table.style = 'Table Grid'
            for line in rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, line): cell.text = '' if value is None else str(value)
        doc.save(buffer)
    elif suffix == '.pdf':
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib import colors
        from xml.sax.saxutils import escape
        with LOCK:
            pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
            style = ParagraphStyle('body', fontName='STSong-Light', fontSize=11, leading=17)
            blocks = []
            for line in content.splitlines():
                blocks.extend([Paragraph(escape(line) or ' ', style), Spacer(1, 4)])
            if rows:
                table = Table([[Paragraph(escape('' if c is None else str(c)), style) for c in row] for row in rows], repeatRows=1)
                table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.4,colors.lightgrey),('VALIGN',(0,0),(-1,-1),'TOP')]))
                blocks.append(table)
            SimpleDocTemplate(buffer, leftMargin=42, rightMargin=42, topMargin=42, bottomMargin=42).build(blocks)
    else:
        raise ValueError('unsupported output')
    return buffer.getvalue()
